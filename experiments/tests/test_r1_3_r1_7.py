"""Tests for the confidence probe (R1-3) and the sign-off agreement script (R1-7)."""
import os
import shutil
import sys

import pytest

EXP = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, EXP)
from r1_3_calibration import calibration_probe as cp  # noqa: E402
from r1_7_agreement import signoff_agreement as sa  # noqa: E402

TEMPLATE = os.path.join(EXP, "r1_7_agreement", "template", "Reference_Plan_SignOff.xlsx")


# --- R1-3 -----------------------------------------------------------------

def test_leave_one_out_excludes_own_outcome():
    ents = [{"k": "a", "y": 1}, {"k": "a", "y": 1}, {"k": "a", "y": 0}, {"k": "b", "y": 0}]
    est = cp.loo(ents, "k")
    assert est[0] == pytest.approx((1 + 1) / (2 + 2))   # others in a: one correct of two
    assert est[2] == pytest.approx((2 + 1) / (2 + 2))
    assert est[3] == pytest.approx(1 / 2)               # no other member: prior 0.5


@pytest.mark.skipif(not os.path.exists(os.path.join(cp.GROUNDTRUTH, "results_v2.json")), reason="ground truth absent")
def test_reference_probe_values():
    r = cp.evaluate(cp.load("reference"))
    assert (r["entities"], r["correct"]) == (22, 15)
    pol = r["predictors"]["policy"]
    assert pol["brier"] == pytest.approx(0.330, abs=0.001)
    assert pol["routing"] == {"reviewed": 3, "wrong_entities": 7, "wrong_routed_to_review": 0, "correct_routed_to_review": 3}
    assert r["predictors"]["base_rate_loo"]["brier"] < pol["brier"]
    s = cp.evaluate(cp.load("swap"))
    assert s["predictors"]["policy"]["routing"]["wrong_routed_to_review"] == 4


# --- R1-7 -----------------------------------------------------------------

def test_kappa_known_values():
    a = ["Correct"] * 20 + ["Wrong"] * 10
    assert sa.cohen_kappa(a, a)["kappa"] == pytest.approx(1.0)
    b = ["Correct"] * 18 + ["Wrong"] * 2 + ["Correct"] * 2 + ["Wrong"] * 8
    k = sa.cohen_kappa(a, b)
    # po = 26/30, pe = (20/30)(20/30) + (10/30)(10/30) = 5/9 -> kappa = 0.7
    assert k["kappa"] == pytest.approx(0.7, abs=1e-4)
    assert sa.cohen_kappa(["Correct"] * 5, ["Correct"] * 5)["kappa"] is None


def _fill(path, verdicts):
    import openpyxl
    wb = openpyxl.load_workbook(path)
    for name, (col, v) in verdicts.items():
        ws = wb[name]
        for r in range(6, ws.max_row + 1):
            if ws.cell(row=r, column=1).value is not None:
                ws.cell(row=r, column=col).value = v(r) if callable(v) else v
    wb.save(path)


@pytest.mark.skipif(not os.path.exists(TEMPLATE), reason="sign-off template absent")
def test_workbook_reading_and_agreement(tmp_path):
    p1, p2 = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    shutil.copy(TEMPLATE, p1)
    shutil.copy(TEMPLATE, p2)
    fill = {"Physical reference": (7, "Agree"), "Reference plan": (9, "Agree"), "Judgements": (3, "Agree"),
            "Step verdicts": (9, "Agree")}
    _fill(p1, fill)
    fill2 = dict(fill)
    fill2["Step verdicts"] = (9, lambda r: "Disagree" if r in (6, 7) else "Agree")  # steps 1 and 2
    _fill(p2, fill2)
    out = sa.main([str(p1), str(p2)])
    r1 = out["reviewers"][0]
    assert r1["overall"]["items"] == 82 and r1["overall"]["Agree"] == 82
    assert r1["step_verdicts_authors_vs_reviewer"]["kappa"] == pytest.approx(1.0)
    r2 = out["reviewers"][1]
    assert r2["sheets"]["Step verdicts"]["Disagree"] == 2
    assert r2["step_verdicts_authors_vs_reviewer"]["percent_agreement"] == pytest.approx(30 / 32)
    assert out["between_reviewers"]["step_verdicts"]["n"] == 32
