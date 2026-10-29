"""Reviewer 1, Comment 7 and Reviewer 2, Comment 5: agreement between the
authors' audit and returned independent sign-off workbooks.

Usage:
    python r1_7_agreement/signoff_agreement.py <returned workbook> [<second returned workbook>]

Reads the four review sheets of Reference_Plan_SignOff.xlsx (physical
reference, reference plan, judgements, step verdicts; 82 items, example rows
excluded) and reports:
  * per sheet and overall: Agree / Agree with change / Disagree / unanswered;
  * for the 32 step verdicts, the reviewer's own binary verdict (Agree and
    Agree with change keep the authors' Correct/Wrong; Disagree inverts it) and
    Cohen's kappa between authors and reviewer, with its large-sample 95 %
    interval; percent agreement is also given because kappa is unstable when
    one category dominates;
  * with two workbooks: Cohen's kappa between the two reviewers on the step
    verdicts and on the three-category verdicts of all 82 items.
The script never fills or alters a workbook.
"""
import json
import math
import sys
from typing import Dict, List, Optional

SHEETS = ("Physical reference", "Reference plan", "Judgements", "Step verdicts")
CATEGORIES = ("Agree", "Agree with change", "Disagree")


def read_workbook(path: str) -> Dict[str, List[Dict]]:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for name in SHEETS:
        ws = wb[name]
        header_row, cols = None, {}
        for r in range(1, 10):
            vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            if "Reviewer verdict" in vals:
                header_row = r
                cols = {v: i + 1 for i, v in enumerate(vals) if v}
                break
        if header_row is None:
            raise ValueError(f"{path}: sheet {name!r} has no 'Reviewer verdict' column")
        auth_col = next((i for v, i in cols.items() if str(v).startswith("Authors' verdict")), None)
        items = []
        for r in range(header_row + 1, ws.max_row + 1):
            key = ws.cell(row=r, column=1).value
            if key is None or str(key).strip() in ("", "0", "P00", "00.00", "J00"):
                continue  # blank or example row
            v = ws.cell(row=r, column=cols["Reviewer verdict"]).value
            items.append({"item": str(key).strip(), "verdict": (str(v).strip() if v else None),
                          "authors": (str(ws.cell(row=r, column=auth_col).value).strip() if auth_col else None)})
        out[name] = items
    return out


def cohen_kappa(a: List[str], b: List[str]) -> Optional[Dict]:
    pairs = [(x, y) for x, y in zip(a, b) if x is not None and y is not None]
    n = len(pairs)
    if n == 0:
        return None
    cats = sorted({x for x, _ in pairs} | {y for _, y in pairs})
    po = sum(1 for x, y in pairs if x == y) / n
    pe = sum((sum(1 for x, _ in pairs if x == c) / n) * (sum(1 for _, y in pairs if y == c) / n) for c in cats)
    if pe >= 1.0:
        return {"n": n, "percent_agreement": po, "kappa": None, "note": "a single category used by both raters"}
    k = (po - pe) / (1 - pe)
    se = math.sqrt(po * (1 - po) / (n * (1 - pe) ** 2))
    return {"n": n, "percent_agreement": round(po, 4), "kappa": round(k, 4),
            "ci95": [round(k - 1.96 * se, 4), round(k + 1.96 * se, 4)]}


def reviewer_step_labels(items: List[Dict]) -> List[Optional[str]]:
    flip = {"Correct": "Wrong", "Wrong": "Correct"}
    out = []
    for it in items:
        if it["verdict"] is None:
            out.append(None)
        elif it["verdict"] == "Disagree":
            out.append(flip.get(it["authors"], None))
        else:
            out.append(it["authors"])
    return out


def summarise(wb: Dict[str, List[Dict]]) -> Dict:
    res = {"sheets": {}}
    tot = {c: 0 for c in CATEGORIES}
    tot["unanswered"] = 0
    for name, items in wb.items():
        c = {k: sum(1 for i in items if i["verdict"] == k) for k in CATEGORIES}
        c["unanswered"] = sum(1 for i in items if i["verdict"] is None)
        c["items"] = len(items)
        res["sheets"][name] = c
        for k in tot:
            tot[k] += c[k]
    tot["items"] = sum(len(v) for v in wb.values())
    res["overall"] = tot
    steps = wb["Step verdicts"]
    res["step_verdicts_authors_vs_reviewer"] = cohen_kappa([s["authors"] for s in steps], reviewer_step_labels(steps))
    return res


def main(paths: List[str]) -> Dict:
    wbs = [read_workbook(p) for p in paths]
    out = {"workbooks": paths, "reviewers": [summarise(w) for w in wbs]}
    if len(wbs) == 2:
        a, b = wbs
        out["between_reviewers"] = {
            "step_verdicts": cohen_kappa(reviewer_step_labels(a["Step verdicts"]), reviewer_step_labels(b["Step verdicts"])),
            "all_items_three_categories": cohen_kappa([i["verdict"] for s in SHEETS for i in a[s]],
                                                      [i["verdict"] for s in SHEETS for i in b[s]]),
        }
    return out


if __name__ == "__main__":
    if not 2 <= len(sys.argv) <= 3:
        raise SystemExit(__doc__)
    print(json.dumps(main(sys.argv[1:]), indent=1))
